"""Собирает таблицу «Об устранении нарушений от ГБУ МАЦ» по образцу эталона.

Структура повторяет эталон: заголовок, две строки шапки с объединёнными
группами ДТ, МКД, ОДХ и ОО, строка нумерации колонок, строки районов и итог.
Второй блок — та же таблица без заполненного округа.

    python work/build_mac_table.py <сводка.csv> <куда положить xlsx> [--месяц АВГУСТ]
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

THIN = Side(style='thin', color='000000')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = 'Arial'
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=False)

GROUPS = [(5, 6), (7, 8), (9, 10), (11, 12)]
SINGLE_COLUMNS = [1, 2, 3, 4, 13, 14, 15, 16]

HEADERS = {
    1: '№ п/п',
    2: 'Округ',
    3: 'Район',
    4: 'Всего нарушений\nза текущий период',
    5: 'Из них по ДТ',
    7: 'Из них по МКД',
    9: 'Из них по ОДХ',
    11: 'Из них по ОО',
    13: 'Всего нарушений\nустранено',
    14: 'Возвращено\nна доработку',
    15: '% устраненных\nнарушений',
    16: 'Остается\nна контроле'
}

SUBMIT = {5: 'Всего', 6: 'Устранено', 7: 'Всего', 8: 'Устранено',
          9: 'Всего', 10: 'Устранено', 11: 'Всего', 12: 'Устранено'}


def read_summary(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.DictReader(handle, delimiter=';'))
    districts = [row for row in rows if row['Район'] != 'Итого']
    return districts


def write_header(sheet, title_row, districts_with_okrug, title_text):
    """Пишет заголовок, шапку в две строки, нумерацию колонок и строки районов."""
    header_row = title_row + 1
    sheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=16)
    title = sheet.cell(row=title_row, column=1, value=title_text)
    title.font = Font(name=FONT, bold=True, size=11)
    title.alignment = CENTER

    for column, text in HEADERS.items():
        cell = sheet.cell(row=header_row, column=column, value=text)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.alignment = CENTER
        cell.border = BORDER
    for column, text in SUBMIT.items():
        cell = sheet.cell(row=header_row + 1, column=column, value=text)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.alignment = CENTER
        cell.border = BORDER
    for column in SINGLE_COLUMNS:
        sheet.merge_cells(start_row=header_row, start_column=column, end_row=header_row + 1, end_column=column)
        sheet.cell(row=header_row + 1, column=column).border = BORDER
    for first, second in GROUPS:
        sheet.merge_cells(start_row=header_row, start_column=first, end_row=header_row, end_column=second)

    number_row = header_row + 2
    for column in range(1, 17):
        cell = sheet.cell(row=number_row, column=column, value=column)
        cell.font = Font(name=FONT, size=9)
        cell.alignment = CENTER
        cell.border = BORDER

    first_data = number_row + 1
    for offset, district in enumerate(districts_with_okrug):
        row = first_data + offset
        sheet.cell(row=row, column=1, value=offset + 1)
        sheet.cell(row=row, column=2, value='САО' if district['Округ'] else None)
        sheet.cell(row=row, column=3, value=district['Район'])
        sheet.cell(row=row, column=4, value=district['Всего'])
        sheet.cell(row=row, column=5, value=district['ДТ'])
        sheet.cell(row=row, column=7, value=district['МКД'])
        sheet.cell(row=row, column=9, value=district['ОДХ'])
        sheet.cell(row=row, column=11, value=district['ОО'])
        sheet.cell(row=row, column=15, value=f'=IFERROR(M{row}/D{row},0)')
        sheet.cell(row=row, column=16, value=f'=D{row}-M{row}')
        for column in range(1, 17):
            cell = sheet.cell(row=row, column=column)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = LEFT if column == 3 else CENTER
        sheet.cell(row=row, column=15).number_format = '0%'

    total_row = first_data + len(districts_with_okrug)
    sheet.cell(row=total_row, column=3, value='Итого:')
    for column in (4, 5, 7, 9, 11, 13, 14):
        letter = get_column_letter(column)
        sheet.cell(row=total_row, column=column, value=f'=SUM({letter}{first_data}:{letter}{total_row - 1})')
    sheet.cell(row=total_row, column=15, value=f'=IFERROR(M{total_row}/D{total_row},0)')
    sheet.cell(row=total_row, column=16, value=f'=D{total_row}-M{total_row}')
    for column in range(1, 17):
        cell = sheet.cell(row=total_row, column=column)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.border = BORDER
        cell.alignment = LEFT if column == 3 else CENTER
    sheet.cell(row=total_row, column=15).number_format = '0%'
    return total_row


def build(districts, month):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = month.capitalize()

    first = [{**district, 'Округ': True} for district in districts]
    second = [{**district, 'Округ': False} for district in districts]

    title = f'Об устранении нарушений от ГБУ «МАЦ» — {month.upper()}'
    first_total = write_header(sheet, 1, first, title)
    second_title = first_total + 2
    write_header(sheet, second_title, second, 'Об устранении нарушений от ГБУ «МАЦ»')

    widths = {1: 7, 2: 9, 3: 24, 4: 15}
    for column in range(5, 17):
        widths[column] = 12
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 20
    sheet.row_dimensions[2].height = 34
    sheet.freeze_panes = 'A5'
    return workbook


def main():
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    month = 'август'
    if '--месяц' in sys.argv:
        month = sys.argv[sys.argv.index('--месяц') + 1]
    districts = read_summary(source)
    workbook = build(districts, month)
    workbook.save(target)
    print(f'таблица собрана: {target}')
    print(f'районов: {len(districts)}, нарушений: {sum(int(row["Всего"]) for row in districts)}')


if __name__ == '__main__':
    main()
