"""Вставляет фотографии нарушений в книгу «Об устранении нарушений от ГБУ МАЦ».

Раскладка повторяет исходные приложения МАЦ: на каждый район отдельный лист,
на каждое нарушение блок из фотографии и строки «Адрес: … Нарушение: …».

    python work/embed_mac_photos.py <книга.xlsx> <папка с фото> [--cap 1600] [--качество 80]
"""
import argparse
import io
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from PIL import Image as Picture

THIN = Side(style='thin', color='000000')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = 'Century Gothic'
FILL = 'EFEFEF'
COLUMNS = 12
COLUMN_WIDTH = 8.86
MAX_PHOTO_WIDTH = 780
MAX_PHOTO_HEIGHT = 520
FORBIDDEN = re.compile(r'[\[\]:*?/\\]')


def read_index(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        item = dict(zip(headers, row))
        if item.get('Район'):
            records.append(item)
    return records


def prepare(photo_path, cache, cap, quality):
    """Готовит уменьшенную копию фотографии и возвращает путь к ней."""
    target = cache / f'{photo_path.stem}.jpg'
    if target.exists():
        return target
    with Picture.open(photo_path) as picture:
        picture = picture.convert('RGB')
        picture.thumbnail((cap, cap), Picture.LANCZOS)
        buffer = io.BytesIO()
        picture.save(buffer, 'JPEG', quality=quality, optimize=True, progressive=True)
    target.write_bytes(buffer.getvalue())
    return target


def sheet_name(title, used):
    name = FORBIDDEN.sub('', title).strip(' .')[:31] or 'Район'
    candidate = name
    counter = 2
    while candidate.lower() in used:
        suffix = f' ({counter})'
        candidate = name[:31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


def build(book, records, source, cache, cap, quality, only=None, log=print):
    used = {name.lower() for name in book.sheetnames}
    districts = []
    for record in records:
        if record['Район'] not in districts:
            districts.append(record['Район'])
    if only:
        districts = [district for district in districts if district in only]

    placed = 0
    for district in districts:
        rows = [item for item in records if item['Район'] == district]
        sheet = book.create_sheet(sheet_name(district, used))
        sheet.sheet_view.showGridLines = False
        for column in range(1, COLUMNS + 1):
            sheet.column_dimensions[get_column_letter(column)].width = COLUMN_WIDTH

        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COLUMNS)
        title = sheet.cell(row=1, column=1, value=f'{district} — фотоматериалы по устранению нарушений, август')
        title.font = Font(name=FONT, bold=True, size=13)
        title.alignment = Alignment(horizontal='center', vertical='center')
        title.fill = _fill()
        title.border = BORDER
        sheet.row_dimensions[1].height = 24

        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COLUMNS)
        subtitle = sheet.cell(row=2, column=1, value=f'Всего нарушений: {len(rows)}')
        subtitle.font = Font(name=FONT, size=11)
        subtitle.alignment = Alignment(horizontal='center', vertical='center')
        subtitle.border = BORDER
        sheet.row_dimensions[2].height = 20
        sheet.row_dimensions[3].height = 8

        cursor = 4
        missed = 0
        for item in rows:
            files = [name.strip() for name in str(item.get('Файл') or '').split('|') if name.strip()]
            photo = None
            for name in files:
                candidate = source / district / name
                if candidate.exists():
                    photo = candidate
                    break
            if photo is None:
                missed += 1
                continue

            prepared = prepare(photo, cache, cap, quality)
            with Picture.open(prepared) as picture:
                width, height = picture.size
            scale = min(MAX_PHOTO_WIDTH / width, MAX_PHOTO_HEIGHT / height, 1.0)
            width, height = max(1, round(width * scale)), max(1, round(height * scale))

            image = ExcelImage(str(prepared))
            image.width = width
            image.height = height
            sheet.add_image(image, f'A{cursor}')
            sheet.row_dimensions[cursor].height = round(height * 0.75) + 4

            text_row = cursor + 1
            sheet.merge_cells(start_row=text_row, start_column=1, end_row=text_row, end_column=COLUMNS)
            text = sheet.cell(row=text_row, column=1, value=record_text(item))
            text.font = Font(name=FONT, size=11)
            text.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            text.border = BORDER
            sheet.row_dimensions[text_row].height = 34
            sheet.row_dimensions[cursor + 2].height = 8
            cursor += 3
            placed += 1
        log(f'{district:<22} нарушений {len(rows):>4}, фото {len(rows) - missed:>4}' + (f', без файла {missed}' if missed else ''))
    return placed


def _fill():
    from openpyxl.styles import PatternFill
    return PatternFill('solid', fgColor=FILL)


def record_text(item):
    address = str(item.get('Адрес') or '').strip()
    violation = str(item.get('Нарушение') or '').strip()
    parts = [f'Адрес: {address}']
    if violation:
        parts.append(f'Нарушение: {violation}')
    return '. '.join(parts)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('book')
    parser.add_argument('photos')
    parser.add_argument('--cap', type=int, default=1600)
    parser.add_argument('--качество', dest='quality', type=int, default=80)
    parser.add_argument('--index', default=None)
    parser.add_argument('--only', default=None, help='районы через запятую, только для проверки')
    parser.add_argument('--out', default=None)
    args = parser.parse_args()

    book_path = Path(args.book)
    source = Path(args.photos)
    index_path = Path(args.index) if args.index else source / 'Указатель.xlsx'
    out_path = Path(args.out) if args.out else book_path
    cache = Path(sys.argv[0]).parent.parent / 'work' / '.photo-cache'
    cache.mkdir(parents=True, exist_ok=True)

    records = read_index(index_path)
    print(f'нарушений в указателе: {len(records)}')
    only = [name.strip() for name in args.only.split(',')] if args.only else None
    book = openpyxl.load_workbook(book_path)
    placed = build(book, records, source, cache, args.cap, args.quality, only)
    book.save(out_path)
    print()
    print(f'вставлено фотографий: {placed}')
    print(f'файл: {out_path} ({out_path.stat().st_size / 1024 / 1024:.0f} МБ)')


if __name__ == '__main__':
    main()
