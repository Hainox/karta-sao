"""Выгружает листы карты ОДХ в отдельную книгу.

Запуск: python work/export_odh_accounts.py <книга со всеми учётками> <новая книга>
"""
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

SOURCE_SHEETS = ['Карта ОДХ', 'Как войти (карта ОДХ)']


def main():
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    book = load_workbook(source)
    missing = [name for name in SOURCE_SHEETS if name not in book.sheetnames]
    if missing:
        raise SystemExit(f'В файле {source.name} нет листов: {", ".join(missing)}')

    result = Workbook()
    result.remove(result.active)
    for name in SOURCE_SHEETS:
        sheet = book[name]
        copy = result.create_sheet(name)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    copy.cell(row=cell.row, column=cell.column, value=cell.value)
        for column, dimension in sheet.column_dimensions.items():
            copy.column_dimensions[column].width = dimension.width
        if sheet.freeze_panes:
            copy.freeze_panes = sheet.freeze_panes
        if sheet.auto_filter and sheet.auto_filter.ref:
            copy.auto_filter.ref = sheet.auto_filter.ref

    result.save(target)
    accounts = result['Карта ОДХ']
    print(f'файл собран: {target}')
    print(f'листов: {len(result.sheetnames)}, учётных записей: {accounts.max_row - 1}')


if __name__ == '__main__':
    main()
